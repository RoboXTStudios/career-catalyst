"""Golden Resume, document hygiene, and submission-readiness regressions."""

from __future__ import annotations

import hashlib
import shutil
import tempfile
import zipfile
from pathlib import Path

import pytest
import yaml
from docx import Document
from lxml import etree

from scripts.docx_quality import (
    compare_docx_factual_parity,
    extract_docx_structure,
    inspect_docx_hygiene,
    sanitize_docx,
    validate_ats_round_trip,
)
from scripts.evidence_engine import load_evidence_cards, select_evidence_cards
from scripts.export_docx import export_ats_docx, export_styled_docx
from scripts.golden_resume import (
    GoldenResumeError,
    load_golden_resume,
    rank_canonical_evidence_cards,
    validate_golden_resume,
)
from scripts.package_generator import PackageGenerationError, generate_package
from scripts.parse_job import parse_job_description
from scripts.submission_readiness import (
    build_interview_conversion_gate,
    build_requirement_coverage_matrix,
    evaluate_claim_provenance,
    evaluate_evidence_density,
    evaluate_human_credibility,
    evaluate_keyword_overuse,
    evaluate_voice_drift,
)
from scripts.tailor_resume import _include_github, render_base_resume
from tests.test_sprint35b_document_writing import _isolated_runtime
from tests.test_sprint39_2_experiential_package_generation import (
    _reopened_live_nation_root,
    _select_three_evidence,
)


ROOT = Path(__file__).resolve().parents[1]
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _export_root(tmp_path: Path) -> tuple[Path, Path]:
    root = tmp_path / "docx-root"
    (root / "data").mkdir(parents=True)
    (root / "templates" / "docx").mkdir(parents=True)
    (root / "exports" / "markdown").mkdir(parents=True)
    shutil.copy2(ROOT / "data" / "platforms.yml", root / "data" / "platforms.yml")
    shutil.copy2(
        ROOT / "templates" / "docx" / "styled_resume_template.docx",
        root / "templates" / "docx" / "styled_resume_template.docx",
    )
    source = root / "exports" / "markdown" / "golden.md"
    source.write_text(render_base_resume(ROOT), encoding="utf-8")
    return root, source


def _traditional_operations_runtime(tmp_path: Path) -> tuple[Path, str]:
    root = tmp_path / "traditional-operations"
    for name in ("data", "config", "templates"):
        shutil.copytree(ROOT / name, root / name)
    (root / "jobs").mkdir()
    fixture = "example_company_senior_operations.md"
    shutil.copy2(ROOT / "tests" / "fixtures" / "jobs" / fixture, root / "jobs" / fixture)
    parsed = parse_job_description(root / "jobs" / fixture)
    selected_ids = [
        "enterprise_media_operations_transformation",
        "operational_workflow_design_airtable_implementation",
        "enterprise_collaboration_platform_adoption_stakeholder_enablement",
    ]
    tracker_id = "sprint42-traditional-operations"
    tracker = {
        "applications": [
            {
                "id": tracker_id,
                "stable_slug": tracker_id,
                "company": parsed["company"],
                "role": parsed["job_title"],
                "status": "Prospect",
                "job_file": f"jobs/{fixture}",
                "priority": "High",
                "show_on_dashboard": True,
                "evidence_project_ids": selected_ids,
                "material_paths": {},
            }
        ]
    }
    (root / "data" / "application_tracker.yml").write_text(
        yaml.safe_dump(tracker, sort_keys=False), encoding="utf-8"
    )
    return root, tracker_id


def _rewrite_zip(path: Path, replacements: dict[str, bytes], additions: dict[str, bytes] | None = None) -> None:
    with tempfile.TemporaryDirectory() as directory:
        target = Path(directory) / path.name
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as destination:
            for info in source.infolist():
                destination.writestr(info, replacements.get(info.filename, source.read(info.filename)))
            for name, payload in (additions or {}).items():
                destination.writestr(name, payload)
        shutil.copy2(target, path)


def test_golden_resume_validates_complete_canonical_inventory():
    inventory = load_golden_resume(ROOT)
    validation = validate_golden_resume(inventory)
    assert validation == {
        "status": "valid",
        "record_count": 40,
        "employment_count": 4,
        "project_count": 4,
        "evidence_project_count": 5,
        "evidence_card_count": 8,
        "skill_group_count": 4,
        "platform_category_count": 5,
        "certification_count": 1,
    }
    assert {project["id"] for project in inventory["projects"]} >= {
        "career_catalyst", "roboxt_studios", "campaignos"
    }
    assert inventory["candidate"]["github_url"].endswith(")")


def test_golden_resume_rejects_stale_provenance_and_preserves_public_employer():
    inventory = load_golden_resume(ROOT)
    assert inventory["employment"][0]["company"] == "OMG23 (Omnicom Media Group)"
    assert "OMD Entertainment" in inventory["employment"][0]["internal_aliases"]
    inventory["evidence_cards"][0]["source"] = "data/achievements.yml:not_real"
    with pytest.raises(GoldenResumeError, match="stale provenance"):
        validate_golden_resume(inventory)


def test_contextual_project_priority_preserves_senior_career_and_role_relevance():
    cards = load_evidence_cards(ROOT)
    product = select_evidence_cards(
        {"job_title": "Director, AI Product Operations", "raw_text": "AI workflow automation product requirements human-in-the-loop builder transformation"},
        cards,
    )
    assert [card["id"] for card in product][:2] == ["career_catalyst", "campaignos"]
    assert "github_product_delivery" in [card["id"] for card in product]
    media = select_evidence_cards(
        {"job_title": "Music Content Operations Lead", "raw_text": "music media creative production content operations digital publishing"},
        cards,
    )
    assert media[0]["id"] == "roboxt_studios"
    ranked = rank_canonical_evidence_cards(cards, "traditional senior operations governance")
    assert ranked[0]["id"] in {"omg23_disney_leadership", "governance_qa_delivery", "martech_campaign_execution"}
    assert _include_github(
        {"job_title": "AI Product Operations Lead", "raw_text": "Build human-in-the-loop AI workflows and product requirements."},
        complete_foundation=False,
    )
    assert not _include_github(
        {"job_title": "Senior Label Relations Manager", "raw_text": "Coordinate music partners and releases."},
        complete_foundation=False,
    )


def test_requirement_matrix_distinguishes_semantic_transfer_and_unsupported_claims():
    inventory = load_golden_resume(ROOT)
    matrix = build_requirement_coverage_matrix(
        {
            "qualifications": [
                "Build a marketing technology ecosystem and workflow governance",
                "Administer Salesforce and write SQL queries",
            ],
            "responsibilities": ["Lead cross-functional program management and executive reporting"],
        },
        inventory,
        [],
        render_base_resume(ROOT),
    )
    by_text = {row["original_jd_wording"]: row for row in matrix}
    assert by_text["Build a marketing technology ecosystem and workflow governance"]["coverage"] in {"PROVEN", "TRANSFERABLE"}
    assert by_text["Administer Salesforce and write SQL queries"]["coverage"] == "NOT_SUPPORTED"
    assert not by_text["Administer Salesforce and write SQL queries"]["evidence_ids"]
    assert by_text["Lead cross-functional program management and executive reporting"]["coverage"] in {"PROVEN", "TRANSFERABLE"}


def test_requirement_matrix_uses_meaningful_requirements_and_preserves_direct_gaps():
    inventory = load_golden_resume(ROOT)
    openai = build_requirement_coverage_matrix(
        parse_job_description(ROOT / "tests/fixtures/sprint35/openai_sales_strategy_operations.md"),
        inventory,
        [],
        render_base_resume(ROOT),
    )
    assert not {"openai", "central", "sales"} & {
        row["original_jd_wording"].lower() for row in openai
    }
    assert next(row for row in openai if "sales forecasting" in row["normalized_requirement"])["coverage"] == "NOT_SUPPORTED"
    assert next(row for row in openai if "territory planning" in row["normalized_requirement"])["coverage"] == "NOT_SUPPORTED"

    twitch = build_requirement_coverage_matrix(
        parse_job_description(ROOT / "tests/fixtures/sprint35/twitch_senior_label_relations_manager.md"),
        inventory,
        [],
        render_base_resume(ROOT),
    )
    label = next(row for row in twitch if "direct label relations" in row["normalized_requirement"])
    assert label["coverage"] == "NOT_SUPPORTED"
    assert not label["evidence_ids"]


def test_voice_specificity_and_human_credibility_are_explainable():
    inventory = load_golden_resume(ROOT)
    vague = "- Drove impactful results and leveraged synergies.\n- Dynamic leader who optimized processes."
    assert evaluate_evidence_density(vague, inventory)["status"] == "REVIEW"
    assert evaluate_voice_drift({"resume": "A dynamic leader with a proven track record."})["status"] == "REVIEW"
    matrix = build_requirement_coverage_matrix(
        {"keywords": ["operations transformation"]}, inventory, [], render_base_resume(ROOT)
    )
    credibility = evaluate_human_credibility(render_base_resume(ROOT), matrix)
    assert credibility["checks"]["professional_identity_visible"]
    assert credibility["checks"]["career_progression_visible"]
    assert credibility["checks"]["unsupported_claims_avoided"]
    repeated = evaluate_keyword_overuse(
        " ".join(["workflow automation ecosystem"] * 5),
        [{"normalized_requirement": "workflow automation ecosystem and governance"}],
    )
    assert repeated["status"] == "PASS"
    stuffed = evaluate_keyword_overuse(
        " ".join(["workflow automation ecosystem and governance"] * 4),
        [{"normalized_requirement": "workflow automation ecosystem and governance"}],
    )
    assert stuffed["status"] == "REVIEW"
    assert evaluate_claim_provenance(render_base_resume(ROOT), inventory)["status"] == "PASS"
    fabricated = render_base_resume(ROOT) + "\n- Administered Salesforce quotas for a global sales organization.\n"
    provenance = evaluate_claim_provenance(fabricated, inventory)
    assert provenance["status"] == "BLOCKED"
    assert any("Salesforce quotas" in claim["text"] for claim in provenance["unresolved_claims"])


def test_docx_hygiene_removes_metadata_comments_and_revision_markup(tmp_path: Path):
    document = Document()
    document.add_paragraph("Final visible candidate text")
    path = tmp_path / "dirty.docx"
    document.save(path)
    with zipfile.ZipFile(path) as archive:
        core = archive.read("docProps/core.xml").replace(
            b"</cp:coreProperties>",
            b"<dc:description>generated by python-docx</dc:description></cp:coreProperties>",
        )
        root = etree.fromstring(archive.read("word/document.xml"))
        run = root.find(f".//{{{W_NS}}}r")
        parent = run.getparent()
        inserted = etree.Element(f"{{{W_NS}}}ins")
        parent.replace(run, inserted)
        inserted.append(run)
        deleted = etree.SubElement(parent, f"{{{W_NS}}}del")
        deleted_run = etree.SubElement(deleted, f"{{{W_NS}}}r")
        etree.SubElement(deleted_run, f"{{{W_NS}}}delText").text = "Deleted draft"
        paragraph = root.find(f".//{{{W_NS}}}p")
        etree.SubElement(paragraph, f"{{{W_NS}}}commentRangeStart").set(f"{{{W_NS}}}id", "0")
        document_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    _rewrite_zip(
        path,
        {"docProps/core.xml": core, "word/document.xml": document_xml},
        {"word/comments.xml": b'<?xml version="1.0"?><w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'},
    )
    before = extract_docx_structure(path)["text"]
    report = sanitize_docx(path)
    assert report["status"] == "PASS"
    assert report["comments"] == report["revisions"] == 0
    assert report["revision_session_attributes"] == 0
    assert not any(report["core_properties"].values())
    assert report["prohibited_generator_identifiers"] == []
    assert report["visible_content_preserved"]
    assert extract_docx_structure(path)["text"] == before == "Final visible candidate text"
    Document(path)


def test_ats_round_trip_and_styled_parity_use_actual_exported_files(tmp_path: Path):
    root, source = _export_root(tmp_path)
    ats = export_ats_docx(source, root)
    styled = export_styled_docx(source, root)
    intended = source.read_text(encoding="utf-8")
    round_trip = validate_ats_round_trip(ats["output_path"], intended)
    parity = compare_docx_factual_parity(styled["output_path"], ats["output_path"])
    assert round_trip["status"] == "PASS"
    assert round_trip["table_count"] == round_trip["textbox_count"] == 0
    assert "Trisha Lynch" in round_trip["parsed_preview"]
    assert "https://github.com/RoboXTStudios" in round_trip["parsed_preview"]
    assert "https://github.com/RoboXTStudios" in round_trip["hyperlinks"]
    assert parity == {
        "status": "PASS",
        "visible_facts_match": True,
        "hyperlinks_match": True,
        "token_delta": {},
    }
    for output in (ats, styled):
        assert inspect_docx_hygiene(output["output_path"])["status"] == "PASS"


def test_interview_gate_surfaces_gaps_without_manufacturing_support():
    coverage = [
        {"coverage": "PROVEN", "normalized_requirement": "program leadership", "original_jd_wording": "Lead programs", "explanation": "Direct support."},
        {"coverage": "NOT_SUPPORTED", "normalized_requirement": "sql expertise", "original_jd_wording": "Expert SQL", "explanation": "No support."},
    ]
    gate = build_interview_conversion_gate(
        parsed_job={"company": "Example", "job_title": "Director"},
        match_report={"match_score": 75},
        coverage_matrix=coverage,
        candidate_qa={"status": "PASS"},
        voice={"status": "PASS"},
        specificity={"status": "PASS", "issues": []},
        credibility={"status": "PASS", "review_items": []},
        ats_round_trip={"status": "PASS"},
        docx_hygiene={"ATS": {"status": "PASS"}, "Styled": {"status": "PASS"}},
        factual_parity={"status": "PASS"},
    )
    assert gate["submission_status"] == "NEEDS REVIEW"
    assert gate["blocking_reasons"] == []
    assert gate["what_career_catalyst_deliberately_did_not_claim"] == ["Expert SQL"]


@pytest.mark.parametrize(
    "role_id,fixture,selected_ids",
    [
        (
            "sprint42-openai",
            "openai_sales_strategy_operations.md",
            ["enterprise_media_operations_transformation", "career_catalyst", "disney_plus_launch_readiness"],
        ),
        (
            "sprint42-twitch",
            "twitch_senior_label_relations_manager.md",
            ["just_for_us_podcast", "roboxt_studios", "enterprise_media_operations_transformation"],
        ),
    ],
)
def test_product_and_media_packages_complete_consolidated_quality_path(
    tmp_path: Path, role_id: str, fixture: str, selected_ids: list[str]
):
    root = _isolated_runtime(tmp_path, role_id, fixture, selected_ids)
    result = generate_package(role_id, root, export_root=tmp_path / "exports")
    quality = result["package_quality"]
    assert result["package_complete"]
    assert quality["ats_round_trip"]["status"] == "PASS"
    assert quality["styled_ats_factual_parity"]["status"] == "PASS"
    assert quality["claim_provenance"]["status"] == "PASS"
    assert all(value["status"] == "PASS" for value in quality["docx_hygiene"].values())
    assert quality["interview_conversion_gate"]["submission_status"] in {"READY TO SUBMIT", "NEEDS REVIEW"}
    files = result["manifest"]["files"]
    for key in ("ats_parsed_preview", "requirement_coverage_matrix", "interview_conversion_gate"):
        assert Path(files[key]).is_file()
    resume_text = Path(files["resume_text"]).read_text(encoding="utf-8")
    ats_structure = extract_docx_structure(files["ats_docx"])
    if "twitch" in role_id:
        assert "https://github.com/RoboXTStudios" not in resume_text


def test_entertainment_package_preserves_selected_evidence_and_transaction(tmp_path: Path, monkeypatch):
    root, tracker_id = _reopened_live_nation_root(tmp_path)
    _select_three_evidence(root, tracker_id)

    def canonical_score(_job, _root, _associated, **_kwargs):
        return {
            "match_score": 62,
            "match_tier": "Stretch Match",
            "match_summary": "Canonical canary score.",
            "match_strengths": ["Transferable production operations."],
            "match_gaps": ["Direct experiential ownership is not claimed."],
            "recommended_action": "Review First",
            "confidence": "Medium",
        }

    monkeypatch.setattr("scripts.package_generator.score_job_match", canonical_score)
    tracker_before = yaml.safe_load((root / "data" / "application_tracker.yml").read_text())
    selected_before = tracker_before["applications"][0]["evidence_project_ids"]
    result = generate_package(tracker_id, root, export_root=tmp_path / "exports")
    saved = yaml.safe_load((root / "data" / "application_tracker.yml").read_text())["applications"][0]
    assert result["match_score"] == 62
    assert saved["evidence_project_ids"] == selected_before
    assert result["manifest"]["prospect_id"] == tracker_id
    assert result["package_quality"]["ats_round_trip"]["status"] == "PASS"


def test_traditional_senior_operations_package_preserves_career_credibility(tmp_path: Path):
    root, tracker_id = _traditional_operations_runtime(tmp_path)
    result = generate_package(tracker_id, root, export_root=tmp_path / "exports")
    quality = result["package_quality"]
    resume = Path(result["manifest"]["files"]["resume_text"]).read_text(encoding="utf-8")
    assert result["package_complete"]
    assert "OMG23 (Omnicom Media Group)" in resume
    assert "Led 10 direct reports" in resume
    assert quality["ats_round_trip"]["status"] == "PASS"
    assert quality["human_credibility"]["checks"]["career_progression_visible"]


def test_docx_hygiene_failure_rolls_back_package_and_tracker(tmp_path: Path, monkeypatch):
    root = _isolated_runtime(
        tmp_path,
        "sprint42-rollback",
        "openai_sales_strategy_operations.md",
        ["enterprise_media_operations_transformation", "career_catalyst", "disney_plus_launch_readiness"],
    )
    export_root = tmp_path / "exports"
    first = generate_package("sprint42-rollback", root, export_root=export_root)
    package = Path(first["saved_package_location"])
    before_hashes = {
        str(path.relative_to(package)): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in package.rglob("*")
        if path.is_file()
    }
    tracker_path = root / "data" / "application_tracker.yml"
    tracker_before = tracker_path.read_bytes()

    def fail_hygiene(_path):
        raise ValueError("injected OOXML hygiene failure")

    monkeypatch.setattr("scripts.export_docx.sanitize_docx", fail_hygiene)
    with pytest.raises(PackageGenerationError, match="Submission-readiness validation blocked"):
        generate_package(
            "sprint42-rollback",
            root,
            force_clean_draft=True,
            export_root=export_root,
        )
    after_hashes = {
        str(path.relative_to(package)): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in package.rglob("*")
        if path.is_file()
    }
    assert after_hashes == before_hashes
    assert tracker_path.read_bytes() == tracker_before
