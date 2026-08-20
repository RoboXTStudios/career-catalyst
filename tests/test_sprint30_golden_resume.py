from __future__ import annotations

import shutil
from pathlib import Path

import pytest
import yaml
from docx import Document

from scripts.export_docx import export_ats_docx, export_styled_docx
from scripts.generate_cover_letter import ApplicationMaterialError, save_material
from scripts.resume_foundation import (
    CandidateLanguageError,
    candidate_language_violations,
    canonical_resume_foundation_info,
    load_resume_foundation,
)
from scripts.tailor_resume import _render_markdown, render_base_resume


PROJECT_ROOT = Path(__file__).resolve().parents[1]

CORRECTED_DISNEY_SCOPE = (
    "Supported Disney Studios, Disney+, Searchlight Pictures, Corporate Brand "
    "Management, and related theatrical, streaming, and franchise priorities."
)

OMG23_PROGRESSION = [
    "Group Director (2022-2026)",
    "Director (2020-2022)",
    "Associate Director (2019-2020)",
    "Supervisor (2018-2019)",
    "Campaign Manager (2016-2018)",
]


def _copy_foundation(tmp_path: Path) -> Path:
    root = tmp_path / "candidate"
    shutil.copytree(PROJECT_ROOT / "data", root / "data")
    shutil.copytree(PROJECT_ROOT / "config", root / "config")
    return root


def test_golden_foundation_preserves_positioning_and_complete_role_history():
    foundation = load_resume_foundation(PROJECT_ROOT)
    data = foundation["data"]
    positions = data["positions"]["positions"]
    omg23 = next(record for record in positions if record["company"].startswith("OMG23"))

    assert data["personal_brand"]["candidate"]["headline"] == (
        "Senior Operations & Transformation Leader | Marketing Technology | AI-Enabled Systems | "
        "Entertainment"
    )
    assert omg23["progression"] == OMG23_PROGRESSION
    assert CORRECTED_DISNEY_SCOPE in omg23["highlights"]
    assert any("10 direct reports" in item for item in omg23["highlights"])
    assert any("64-person organization" in item for item in omg23["highlights"])
    assert any(
        record["progression"]
        == ["Senior Digital Media Planner & Strategist | Media Operations, Search & Social"]
        for record in positions
    )
    assert any(
        record["progression"]
        == ["Senior Digital Media Planner | Strategy, Media Operations & Analytics"]
        for record in positions
    )


def test_real_base_generator_preserves_products_tools_and_golden_language():
    resume = render_base_resume(PROJECT_ROOT)

    for expected in [
        "Senior Operations & Transformation Leader",
        *OMG23_PROGRESSION,
        "Media Operations, Search & Social",
        "Strategy, Media Operations & Analytics",
        "Career Catalyst",
        "RoboXT Studios",
        "CampaignOS",
        "Working Prototype",
        "Google Marketing Platform",
        "Microsoft 365",
        "Workflow Design",
        "ChatGPT",
        "WordPress",
    ]:
        assert expected in resume
    assert CORRECTED_DISNEY_SCOPE in resume
    assert candidate_language_violations(resume) == []


def test_real_docx_exporters_accept_complete_golden_platform_taxonomy(
    tmp_path: Path,
):
    root = tmp_path / "export"
    (root / "data").mkdir(parents=True)
    (root / "templates/docx").mkdir(parents=True)
    (root / "exports/markdown").mkdir(parents=True)
    shutil.copy2(PROJECT_ROOT / "data/platforms.yml", root / "data/platforms.yml")
    shutil.copy2(
        PROJECT_ROOT / "templates/docx/styled_resume_template.docx",
        root / "templates/docx/styled_resume_template.docx",
    )
    source = root / "exports/markdown/golden_master.md"
    source.write_text(render_base_resume(PROJECT_ROOT), encoding="utf-8")

    ats = export_ats_docx(source, root)
    styled = export_styled_docx(source, root)
    ats_document = Document(ats["output_path"])
    styled_document = Document(styled["output_path"])

    assert ats_document.tables == []
    assert len(styled_document.tables[0].columns) == 3
    assert len(styled_document.tables[0].rows) == 2
    table_text = "\n".join(
        cell.text
        for row in styled_document.tables[0].rows
        for cell in row.cells
    )
    assert "Publishing & Creative" in table_text


def test_tailored_generator_inherits_golden_baseline_and_policy():
    foundation = load_resume_foundation(PROJECT_ROOT)
    parsed_job = {
        "job_title": "Senior Operations Leader",
        "company": "Example Company",
        "keywords": ["operations", "transformation", "governance"],
        "raw_text": "Lead operating-model transformation and workflow governance.",
    }
    resume = _render_markdown(
        foundation,
        parsed_job,
        {"keyword_matches": [], "transferable_strengths": []},
        "executive_operations",
    )

    assert "Senior Operations & Transformation Leader" in resume
    assert candidate_language_violations(resume) == []


def test_active_candidate_source_contains_no_unsupported_or_age_signaling_claims():
    info = canonical_resume_foundation_info(PROJECT_ROOT)
    foundation = load_resume_foundation(PROJECT_ROOT)
    active_text = "\n".join(
        str(foundation["data"][Path(relative_path).stem])
        for relative_path in info["baseline_files"]
        if Path(relative_path).stem != "personal_brand"
    ) + "\n" + str({
        key: value
        for key, value in foundation["data"]["personal_brand"].items()
        if key != "claim_guardrails"
    })

    assert CORRECTED_DISNEY_SCOPE in active_text
    assert candidate_language_violations(active_text) == []


def test_existing_evidence_ids_and_provenance_are_preserved():
    evidence = yaml.safe_load(
        (PROJECT_ROOT / "data/evidence_projects.yml").read_text(encoding="utf-8")
    )["evidence_projects"]
    cards = yaml.safe_load(
        (PROJECT_ROOT / "config/evidence_cards.yml").read_text(encoding="utf-8")
    )["evidence_cards"]

    assert {record["id"] for record in evidence} >= {
        "enterprise_media_operations_transformation",
        "enterprise_collaboration_platform_adoption_stakeholder_enablement",
        "operational_workflow_design_airtable_implementation",
        "enterprise_employee_engagement_community_fundraising_initiative",
        "multiverse_editorial",
        "disney_plus_launch_readiness",
        "career_catalyst",
        "roboxt_studios",
        "campaignos",
        "just_for_us_podcast",
    }
    assert {(record["id"], record["source"]) for record in cards} == {
        (
            "omg23_disney_leadership",
            "data/achievements.yml:omg23_advancement,cross_functional_leadership",
        ),
        (
            "campaignos",
            "data/projects.yml:CampaignOS;data/achievements.yml:campaignos_ai_operations",
        ),
        ("career_catalyst", "data/projects.yml:Career Catalyst"),
        ("github_product_delivery", "data/projects.yml:Career Catalyst"),
        ("roboxt_studios", "data/projects.yml:RoboXT Studios"),
        (
            "governance_qa_delivery",
            "data/achievements.yml:workflow_governance,disney_plus_launch_support",
        ),
        (
            "martech_campaign_execution",
            "data/achievements.yml:google_youtube_platform_familiarity",
        ),
        ("photography_creative_voice", "personal creative practice"),
    }
    card_by_id = {record["id"]: record for record in cards}
    assert "10 direct reports" in " ".join(
        card_by_id["omg23_disney_leadership"]["proof_points"]
    )
    assert "working prototype" in card_by_id["campaignos"]["short_description"].lower()
    assert "launched public creative platform" in (
        card_by_id["roboxt_studios"]["short_description"].lower()
    )


def test_excluded_evidence_keeps_provenance_without_becoming_candidate_claim(
    tmp_path: Path,
):
    root = _copy_foundation(tmp_path)
    evidence_path = root / "data/evidence_projects.yml"
    payload = yaml.safe_load(evidence_path.read_text(encoding="utf-8"))
    payload["evidence_projects"].append(
        {
            "id": "historical_private_record",
            "title": "Historical private record",
            "status": "Active",
            "usage_control": "exclude",
            "provenance": {"source": "preserved historical record"},
            "summary": "Supported National Geographic.",
        }
    )
    evidence_path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")

    loaded = load_resume_foundation(root)
    preserved = loaded["data"]["evidence_projects"]["evidence_projects"][-1]
    assert preserved["id"] == "historical_private_record"
    assert preserved["provenance"]["source"] == "preserved historical record"

    payload["evidence_projects"][-1]["usage_control"] = "external"
    evidence_path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    with pytest.raises(CandidateLanguageError, match="National Geographic"):
        load_resume_foundation(root)


def test_legitimate_foundation_change_is_seen_without_stale_cache(tmp_path: Path):
    root = _copy_foundation(tmp_path)
    first = load_resume_foundation(root)
    personal_brand_path = root / "data/personal_brand.yml"
    payload = yaml.safe_load(personal_brand_path.read_text(encoding="utf-8"))
    payload["candidate"]["headline"] = "Operations Transformation Leader"
    personal_brand_path.write_text(
        yaml.safe_dump(payload, sort_keys=False),
        encoding="utf-8",
    )
    second = load_resume_foundation(root)

    assert first["data"]["personal_brand"]["candidate"]["headline"] != (
        second["data"]["personal_brand"]["candidate"]["headline"]
    )
    assert second["data"]["personal_brand"]["candidate"]["headline"] == (
        "Operations Transformation Leader"
    )


@pytest.mark.parametrize(
    "unsupported",
    [
        "Supported FX priorities.",
        "A seasoned operations leader.",
        "Earned an MBA.",
    ],
)
def test_material_save_rejects_unsupported_candidate_claims(
    tmp_path: Path,
    unsupported: str,
):
    context = {
        "root": tmp_path,
        "application": {"company": "Example", "role": "Operations Leader"},
        "parsed_job": {
            "company": "Example",
            "job_title": "Operations Leader",
            "raw_text": "Lead operations and workflow governance.",
        },
        "career_data": {
            "data": {"personal_brand": {"candidate": {"name": "Trisha Lynch"}}}
        },
        "match_report": {"match_score": 90},
        "voice": {"avoid": []},
        "writing_voice": {"banned_phrases": []},
        "material_editing_plan": {"banned_phrases": []},
    }
    content = "\n\n".join([unsupported] * 12)

    with pytest.raises(ApplicationMaterialError, match="Golden Master"):
        save_material(context, "Cover_Letter", content, 5, 500)
