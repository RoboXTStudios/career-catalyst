import shutil
import zipfile
from pathlib import Path

import pytest
import yaml
from docx import Document

from scripts.docx_metadata import clean_docx_metadata
from scripts.export_docx import export_ats_docx, export_styled_docx
from scripts.generate_cover_letter import (
    _dynamic_cover_letter_content,
    _export_cover_letter_docx,
)
from scripts.human_positioning import validate_applicant_evidence
from scripts.public_advocacy import rewrite_public_advocacy
from scripts.tailor_resume import RESUME_EVIDENCE_LIMIT, tailor_resume


ROOT = Path(__file__).resolve().parents[1]
SELECTED_IDS = (
    "campaign_operations_leadership",
    "workflow_design",
    "cross_functional_technical_translation",
    "people_leadership",
)
CLAIM_SAFETY_DISCLAIMERS = (
    "while staying precise about the scope of my direct experience",
    "based on the evidence available",
    "where my experience directly aligns",
    "keep claims close to the facts",
)


def _project_root(tmp_path: Path) -> Path:
    shutil.copytree(ROOT / "config", tmp_path / "config")
    shutil.copytree(ROOT / "data", tmp_path / "data")
    (tmp_path / "jobs").mkdir()
    (tmp_path / "jobs" / "marketing_operations.md").write_text(
        """# Director, Marketing Operations

Company: Example Company
Location: Remote

## Job Description

Lead annual planning, prioritization, process design, executive communication, and cross-functional leadership.
Develop talent through coaching and team leadership. Own Salesforce, Marketo, CRM, Revenue Operations, and B2B SaaS systems.
""",
        encoding="utf-8",
    )
    return tmp_path


def _manual_selection(root: Path) -> dict:
    profile = yaml.safe_load(
        (root / "data" / "evidence_profile.yml").read_text(encoding="utf-8")
    )
    by_id = {item["id"]: item for item in profile["evidence"]}
    selected = []
    for evidence_id in SELECTED_IDS:
        item = dict(by_id[evidence_id])
        item["selected_by"] = "User override"
        item["selection_level"] = "Primary"
        selected.append(item)
    return {
        "primary_evidence": selected,
        "supporting_evidence": [],
        "selected_evidence_ids": list(SELECTED_IDS),
        "overrides": {
            "included_ids": list(SELECTED_IDS),
            "primary_ids": list(SELECTED_IDS),
            "user_reviewed": True,
        },
    }


@pytest.fixture
def generated_resume(tmp_path, monkeypatch):
    root = _project_root(tmp_path)
    selection = _manual_selection(root)

    def score(*_args, **_kwargs):
        return {
            "match_score": 88,
            "match_band": "Strong",
            "match_tier": "Strong Match",
            "top_matching_skills": [],
            "top_matching_projects": [],
            "top_matching_experience": [],
            "tailoring_notes": [],
            "role_interpretation": {},
            "role_evidence_selection": selection,
            "evidence_gap_analysis": {},
        }

    monkeypatch.setattr("scripts.tailor_resume.score_job_match", score)
    result = tailor_resume(
        "executive_operations",
        "jobs/marketing_operations.md",
        root,
        {"role_evidence_selection": selection},
    )
    content = Path(result["output_path"]).read_text(encoding="utf-8")
    return root, result, content, selection


def _document_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _property_xml(path: Path) -> bytes:
    with zipfile.ZipFile(path) as package:
        return b"\n".join(
            package.read(name)
            for name in package.namelist()
            if name.startswith("docProps/") and name.endswith(".xml")
        )


def test_four_manual_evidence_items_reach_resume_and_remain_authoritative(
    generated_resume,
):
    _root, result, content, selection = generated_resume
    selected = result["resume_selected_evidence"]

    assert RESUME_EVIDENCE_LIMIT == 4
    assert result["resume_evidence_limit"] == 4
    assert [item["id"] for item in selected] == list(SELECTED_IDS)
    assert [item["id"] for item in selected] == selection["selected_evidence_ids"]
    for item in selected:
        assert item["description"] in content


def test_supported_signals_and_talent_development_survive_without_inventing_jd_terms(
    generated_resume,
):
    _root, result, content, _selection = generated_resume
    lowered = content.lower()

    assert "people leadership" in lowered
    assert "coaching" in lowered
    assert "led supervisors" in lowered
    assert "led enterprise campaign operations" in lowered
    assert "translated technical platform requirements" in lowered
    assert not any(
        term in lowered
        for term in ("salesforce", "marketo", "revenue operations", "b2b saas")
    )
    assert all(
        item["description"].split()[0] not in {"Supported", "Assisted", "Coordinated"}
        for item in result["resume_selected_evidence"]
    )


def test_ats_and_styled_resumes_receive_all_four_selected_evidence_items(
    generated_resume,
):
    root, result, _content, _selection = generated_resume
    expected = [item["description"] for item in result["resume_selected_evidence"]]

    for exporter in (export_ats_docx, export_styled_docx):
        exported = exporter(result["output_path"], root)
        rendered = _document_text(Path(exported["output_path"]))
        assert all(value in rendered for value in expected)


def test_all_generated_docx_types_use_candidate_metadata_and_no_python(
    generated_resume,
):
    root, result, _content, _selection = generated_resume
    paths = [
        Path(export_ats_docx(result["output_path"], root)["output_path"]),
        Path(export_styled_docx(result["output_path"], root)["output_path"]),
    ]
    cover_path = root / "cover_letter.md"
    paths.append(
        _export_cover_letter_docx(
            {
                "parsed_job": {
                    "job_title": "Director, Marketing Operations",
                    "company": "Example Company",
                }
            },
            "Hello,\n\nA concise, grounded letter.\n\nBest,\n\nTrisha Lynch",
            cover_path,
        )
    )

    for path in paths:
        properties = Document(path).core_properties
        assert properties.author == "Trisha Lynch"
        assert properties.last_modified_by == "Trisha Lynch"
        assert "python" not in _property_xml(path).decode("utf-8").lower()
        assert "generated by python" not in _property_xml(path).decode("utf-8").lower()


def test_metadata_cleanup_changes_properties_without_changing_document_text(tmp_path):
    path = tmp_path / "metadata.docx"
    document = Document()
    document.add_paragraph("Body text remains byte-for-byte meaningful: Python belongs here.")
    document.core_properties.author = "Python"
    document.core_properties.comments = "Generated by python-docx"
    document.save(path)
    before = _document_text(path)

    clean_docx_metadata(path, title="Trisha Lynch - Resume")

    assert _document_text(path) == before
    assert "python" not in _property_xml(path).decode("utf-8").lower()


def test_cover_letter_removes_claim_safety_disclaimers_and_keeps_claim_validation():
    injected = (
        "Hello,\n\nI led cross-functional campaign operations. Based on the evidence available, "
        "where my experience directly aligns, I can contribute.\n\nBest,\n\nTrisha Lynch"
    )
    rewritten, review = rewrite_public_advocacy(
        injected, company="Example Company", role="Director, Marketing Operations"
    )
    lowered = rewritten.lower()
    assert review["valid"] is True
    assert not any(phrase in lowered for phrase in CLAIM_SAFETY_DISCLAIMERS)

    context = {
        "parsed_job": {
            "company": "Example Company",
            "job_title": "Director, Marketing Operations",
        },
        "effective_voice_profile": {"role_family": "business_operations"},
        "selected_evidence_cards": [],
    }
    generated = _dynamic_cover_letter_content(context).lower()
    assert not any(phrase in generated for phrase in CLAIM_SAFETY_DISCLAIMERS)

    with pytest.raises(ValueError):
        validate_applicant_evidence(
            "I built an unsupported personal Career Catalyst project.",
            "cover letter",
        )


def test_dynamic_cover_letter_uses_only_two_strongest_selected_evidence_stories():
    context = {
        "parsed_job": {
            "company": "Example Company",
            "job_title": "Director, Marketing Operations",
        },
        "effective_voice_profile": {"role_family": "business_operations"},
        "selected_evidence_cards": [
            {"id": "one", "proof_points": ["Led enterprise workflow design."]},
            {"id": "two", "proof_points": ["Built executive decision systems."]},
            {"id": "three", "proof_points": ["Coordinated an unrelated task."]},
        ],
    }

    content = _dynamic_cover_letter_content(context)
    assert "I led enterprise workflow design." in content
    assert "I built executive decision systems." in content
    assert "Coordinated an unrelated task" not in content
