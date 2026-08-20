from __future__ import annotations

import shutil
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from scripts.cli import main
from scripts.docx_quality import inspect_docx_hygiene
from scripts.export_docx import MissingMarkdownFileError, export_ats_docx, export_styled_docx


PROJECT_ROOT = Path(__file__).resolve().parents[1]
LINKEDIN_URL = "https://www.linkedin.com/in/trisha-lynch-3433417"
GITHUB_URL = "https://github.com/RoboXTStudios"


def _root(tmp_path: Path) -> tuple[Path, Path]:
    root = tmp_path / "runtime"
    for directory in ("config", "data", "templates"):
        shutil.copytree(PROJECT_ROOT / directory, root / directory)
    markdown = root / "input" / "resume.md"
    markdown.parent.mkdir(parents=True)
    markdown.write_text(
        """<!-- career-catalyst-job-title: Director, Enterprise Strategy -->
<!-- career-catalyst-company: Crunchyroll -->

# Trisha Lynch

Senior Operations & Transformation Leader

Los Angeles, CA | tslynch@mac.com | LinkedIn: [https://www.linkedin.com/in/trisha-lynch-3433417](https://www.linkedin.com/in/trisha-lynch-3433417) | GitHub: [https://github.com/RoboXTStudios](https://github.com/RoboXTStudios)

## Profile

Operations leader who turns complex priorities into governed workflows and dependable execution.

## Core Competencies

- Enterprise Strategy
- Workflow Governance
- AI Workflow Design

## Platforms & Technologies

### Marketing Technology & Measurement

Campaign Manager 360, Google Analytics

### Business Productivity & Collaboration

Microsoft 365, Airtable

### Operations & Program Management

Workflow Design, Process Automation

### AI, Automation & Product Development

Career Catalyst, Python (Working Knowledge)

### Publishing & Creative

Newsletter Development, Editorial Production

## Professional Experience

### OMG23 / OMD Entertainment, Omnicom Media Group

Group Director | 2022-2026

- Led 10 direct reports and provided strategic and operational leadership across an integrated 64-person organization.

## Relevant Projects & Impact

### Career Catalyst

- Built an AI-enabled career intelligence and application-operations product.

### CampaignOS (Working Prototype)

- Built a working prototype that standardizes workflow governance and readiness validation.
""",
        encoding="utf-8",
    )
    return root, markdown


def _document_text(document: Document) -> str:
    body = [paragraph.text for paragraph in document.paragraphs]
    tables = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(body + tables)


def test_exports_create_non_empty_current_filenames(tmp_path: Path):
    root, markdown = _root(tmp_path)
    styled = export_styled_docx(markdown, root)
    ats = export_ats_docx(markdown, root)
    assert Path(styled["output_path"]).name == "crunchyroll_director_enterprise_strategy_trisha_lynch_styled_resume.docx"
    assert Path(ats["output_path"]).name == "crunchyroll_director_enterprise_strategy_trisha_lynch_ats_resume.docx"
    assert all(Path(item["output_path"]).stat().st_size > 0 for item in (styled, ats))


def test_missing_markdown_file_produces_helpful_error(tmp_path: Path):
    root, _markdown = _root(tmp_path)
    with pytest.raises(MissingMarkdownFileError, match="Markdown resume file not found"):
        export_styled_docx(root / "missing.md", root)


def test_ats_is_single_column_and_styled_retains_layout(tmp_path: Path):
    root, markdown = _root(tmp_path)
    styled = Document(export_styled_docx(markdown, root)["output_path"])
    ats = Document(export_ats_docx(markdown, root)["output_path"])
    assert styled.paragraphs
    assert styled.tables
    assert not ats.tables
    for section in ats.sections:
        columns = section._sectPr.find(qn("w:cols"))
        assert columns is not None and columns.get(qn("w:num")) == "1"


def test_exports_preserve_canonical_visible_content_and_links(tmp_path: Path):
    root, markdown = _root(tmp_path)
    for exporter in (export_styled_docx, export_ats_docx):
        result = exporter(markdown, root)
        document = Document(result["output_path"])
        content = _document_text(document)
        for expected in (
            "OMG23 / OMD Entertainment, Omnicom Media Group",
            "Career Catalyst",
            "CampaignOS (Working Prototype)",
            "AI Workflow Design",
            "Process Automation",
            "Python (Working Knowledge)",
            LINKEDIN_URL,
            GITHUB_URL,
        ):
            assert expected in content
        targets = {
            relationship.target_ref
            for relationship in document.part.rels.values()
            if relationship.reltype == RT.HYPERLINK
        }
        assert {LINKEDIN_URL, GITHUB_URL} <= targets


def test_exports_have_clean_ooxml_and_ats_round_trip(tmp_path: Path):
    root, markdown = _root(tmp_path)
    styled = export_styled_docx(markdown, root)
    ats = export_ats_docx(markdown, root)
    for result in (styled, ats):
        hygiene = inspect_docx_hygiene(result["output_path"])
        assert hygiene["comments"] == 0
        assert hygiene["revisions"] == 0
        assert hygiene["prohibited_generator_identifiers"] == []
    assert ats["ats_round_trip"]["status"] == "PASS"


def test_styled_export_uses_template_when_present(tmp_path: Path):
    root, markdown = _root(tmp_path)
    result = export_styled_docx(markdown, root)
    assert result["template_path"] == str(root / "templates" / "docx" / "styled_resume_template.docx")


def test_backward_compatible_cli_defaults_to_styled(tmp_path: Path, monkeypatch):
    root, markdown = _root(tmp_path)
    monkeypatch.chdir(root)
    monkeypatch.setattr("scripts.cli.PROJECT_ROOT", root)
    output = StringIO()
    with redirect_stdout(output):
        exit_code = main(["export-docx", str(markdown)])
    assert exit_code == 0
    assert "Export mode: styled" in output.getvalue()
    assert (root / "exports" / "docx" / "crunchyroll_director_enterprise_strategy_trisha_lynch_styled_resume.docx").is_file()
