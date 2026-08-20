import re
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.export_docx import export_ats_docx, export_styled_docx
from scripts.generate_application_note import generate_application_note
from scripts.generate_cover_letter import generate_cover_letter
from scripts.generate_messages import generate_message
from scripts.generate_strategy_pack import generate_strategy_pack
from scripts.tailor_resume import tailor_resume
from scripts.text_cleanup import cleanup_repeated_words


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PLAYSTATION_JOB = "jobs/playstation_head_global_creative_product_dev_ops.md"
REPEATED_WORD_PATTERN = re.compile(
    r"\b(role|this|the|opportunity|that)\s+\1\b",
    re.IGNORECASE,
)


def _document_text(document):
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(paragraphs + table_cells)


class EntertainmentCalibrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temporary.name)
        for directory in ("data", "config", "templates", "jobs"):
            shutil.copytree(PROJECT_ROOT / directory, cls.root / directory)
        cls.results = {
            "resume": tailor_resume("executive_operations", PLAYSTATION_JOB, cls.root),
            "cover_letter": generate_cover_letter(PLAYSTATION_JOB, cls.root),
            "recruiter": generate_message("recruiter", PLAYSTATION_JOB, cls.root),
            "hiring_manager": generate_message(
                "hiring-manager", PLAYSTATION_JOB, cls.root
            ),
            "application_note": generate_application_note(PLAYSTATION_JOB, cls.root),
            "strategy_pack": generate_strategy_pack(PLAYSTATION_JOB, cls.root),
        }
        cls.contents = {
            name: Path(result["output_path"]).read_text(encoding="utf-8")
            for name, result in cls.results.items()
        }

        markdown_path = cls.results["resume"]["output_path"]
        cls.styled_result = export_styled_docx(markdown_path, cls.root)
        cls.ats_result = export_ats_docx(markdown_path, cls.root)
        cls.styled_document = Document(cls.styled_result["output_path"])
        cls.ats_document = Document(cls.ats_result["output_path"])

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def test_cover_letter_opening_is_clean(self):
        cover_letter = self.contents["cover_letter"]

        self.assertTrue(
            cover_letter.startswith(
                "Dear Sony Interactive Entertainment / PlayStation Hiring Team,\n\n"
            )
        )
        self.assertNotIn("role role", cover_letter.lower())
        self.assertNotIn("PlayStation's", cover_letter)

    def test_cover_letter_uses_employer_shorthand_after_full_mention(self):
        cover_letter = self.contents["cover_letter"]
        full_name = "OMG23 / OMD Entertainment, Omnicom Media Group"

        self.assertEqual(cover_letter.count(full_name), 1)

    def test_playstation_outputs_use_upload_friendly_filenames(self):
        expected_suffixes = {
            "cover_letter": "_cover_letter.txt",
            "recruiter": "_recruiter_message.txt",
            "hiring_manager": "_hiring_manager_message.txt",
            "application_note": "_application_note.txt",
            "strategy_pack": "_strategy_pack.txt",
        }
        for result_name, expected_suffix in expected_suffixes.items():
            filename = Path(self.results[result_name]["output_path"]).name
            self.assertTrue(filename.endswith(expected_suffix), filename)
            self.assertEqual(filename, filename.lower())

        self.assertEqual(
            Path(self.styled_result["output_path"]).name,
            "playstation_head_global_creative_product_development_operations_trisha_lynch_styled_resume.docx",
        )
        self.assertEqual(
            Path(self.ats_result["output_path"]).name,
            "playstation_head_global_creative_product_development_operations_trisha_lynch_ats_resume.docx",
        )

    def test_playstation_cover_letter_has_plain_text_export(self):
        markdown_path = Path(self.results["cover_letter"]["output_path"])
        text_path = Path(self.results["cover_letter"]["txt_output_path"])

        self.assertEqual(
            text_path.name,
            "playstation_head_global_creative_product_development_operations_trisha_lynch_cover_letter.txt",
        )
        self.assertTrue(text_path.is_file())
        self.assertEqual(
            markdown_path.read_text(encoding="utf-8"),
            text_path.read_text(encoding="utf-8"),
        )

    def test_upload_filenames_have_no_unsafe_characters(self):
        output_paths = [Path(result["output_path"]) for result in self.results.values()]
        output_paths.extend(
            [
                Path(self.results["cover_letter"]["txt_output_path"]),
                Path(self.styled_result["output_path"]),
                Path(self.ats_result["output_path"]),
            ]
        )
        for output_path in output_paths:
            self.assertIsNone(re.search(r"[/,()&\s]", output_path.name))

    def test_cover_letter_experience_paragraph_is_concise(self):
        cover_letter = self.contents["cover_letter"]
        experience_paragraph = cover_letter.split("\n\n")[2]
        word_count = len(
            re.findall(
                r"\b[A-Za-z0-9]+(?:['-][A-Za-z0-9]+)*\b",
                experience_paragraph,
            )
        )

        self.assertLess(word_count, 130)
        self.assertLessEqual(experience_paragraph.lower().count("campaign operations"), 1)
        self.assertLessEqual(experience_paragraph.lower().count("complex"), 1)
        self.assertLessEqual(experience_paragraph.lower().count("cross-functional"), 1)
        self.assertIn("OMG23 / OMD Entertainment, Omnicom Media Group", experience_paragraph)
        self.assertIn("10 direct reports", experience_paragraph)
        self.assertIn("64-person organization", experience_paragraph)

    def test_generated_materials_have_no_obvious_repeated_words(self):
        for name, content in self.contents.items():
            self.assertIsNone(
                REPEATED_WORD_PATTERN.search(content),
                f"Repeated adjacent word in {name}",
            )

    def test_cleanup_helper_preserves_unlisted_repetition(self):
        cleaned = cleanup_repeated_words("Duran Duran role role and the the plan")

        self.assertEqual(cleaned, "Duran Duran role and the plan")

    def test_playstation_materials_preserve_supported_entertainment_context(self):
        combined = "\n".join(self.contents.values())
        self.assertRegex(combined.lower(), r"theatrical|streaming")
        for name, content in self.contents.items():
            self.assertNotIn("FX", content, name)
            self.assertNotIn("Disney Branded Television", content, name)
            self.assertNotIn("National Geographic", content, name)

    def test_playstation_materials_prioritize_creative_and_product_operations(self):
        combined = "\n".join(self.contents.values()).lower()
        self.assertIn("creative", combined)
        self.assertIn("product", combined)
        self.assertIn("operations", combined)
        self.assertNotIn("networks", combined)

    def test_campaignos_is_contextual_not_mandatory_in_every_material(self):
        for name, content in self.contents.items():
            if "CampaignOS" in content:
                self.assertIn("prototype", content.lower(), name)

    def test_playstation_docx_exports_preserve_taxonomy_and_table_invariants(self):
        styled_text = _document_text(self.styled_document)
        ats_text = _document_text(self.ats_document)

        self.assertGreaterEqual(len(self.styled_document.tables), 1)
        self.assertEqual(len(self.ats_document.tables), 0)
        for content in (styled_text, ats_text):
            self.assertIn("OMG23 / OMD Entertainment, Omnicom Media Group", content)
            self.assertIn("10 direct reports", content)
            self.assertIn("64-person organization", content)

    def test_all_playstation_outputs_exist_and_are_non_empty(self):
        output_paths = [Path(result["output_path"]) for result in self.results.values()]
        output_paths.extend(
            [
                Path(self.results["cover_letter"]["txt_output_path"]),
                Path(self.styled_result["output_path"]),
                Path(self.ats_result["output_path"]),
            ]
        )
        for output_path in output_paths:
            self.assertTrue(output_path.is_file())
            self.assertGreater(output_path.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
