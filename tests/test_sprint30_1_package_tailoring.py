from __future__ import annotations

import json
import shutil
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document

from scripts.career_claims import (
    historical_omg23_aliases,
    public_claim_violations,
)
from scripts.filename_utils import company_display_name
from scripts.export_docx import export_ats_docx, export_styled_docx
from scripts.generate_cover_letter import (
    _cover_letter_content,
    _word_count,
    load_generation_context,
)
from scripts.materials_library import organize_package_outputs
from scripts.package_generator import preflight_package_generation
from scripts.resume_foundation import load_resume_foundation
from scripts.tailor_resume import _render_markdown


PROJECT_ROOT = Path(__file__).resolve().parents[1]
GITLAB_FIXTURE = Path("tests/fixtures/jobs/gitlab_ai_transformation_owner_marketing.md")


def _application(**changes):
    record = {
        "id": "gitlab_ai_transformation_owner_marketing",
        "stable_slug": "gitlab_ai_transformation_owner_marketing",
        "company": "Gitlab",
        "role": "AI Transformation Owner, Marketing",
        "status": "Drafted",
        "job_id": "8548105002",
        "job_file": str(GITLAB_FIXTURE),
        "material_paths": {},
    }
    record.update(changes)
    return record


def _parsed_gitlab_job():
    return {
        "company": "GitLab",
        "job_title": "AI Transformation Owner, Marketing",
        "keywords": [
            "AI transformation", "Marketing operations", "adoption", "change management",
            "workflow", "governance", "schemas", "validation", "Champion network",
        ],
        "raw_text": GITLAB_FIXTURE.read_text(encoding="utf-8"),
    }


def test_job_description_outside_export_root_is_not_legacy_material(tmp_path):
    job = tmp_path / "jobs/gitlab.md"
    job.parent.mkdir(parents=True)
    job.write_text("job", encoding="utf-8")
    record = _application(material_paths={"Job Description": str(job)})

    result = preflight_package_generation(
        record["id"], [record], PROJECT_ROOT, export_root=tmp_path / "exports"
    )

    assert result["status"] == "ready"
    assert result["conflicts"] == []
    assert result["blocking_issues"] == []


def test_generated_material_outside_export_root_remains_a_conflict(tmp_path):
    resume = tmp_path / "legacy/resume.docx"
    resume.parent.mkdir(parents=True)
    resume.write_bytes(b"resume")
    record = _application(material_paths={"ATS Resume": str(resume)})

    result = preflight_package_generation(
        record["id"], [record], PROJECT_ROOT, export_root=tmp_path / "exports"
    )

    assert result["status"] == "conflict"
    assert result["conflicts"][0]["path"] == str(resume.resolve())


def test_material_owned_by_another_prospect_remains_blocked(tmp_path):
    shared = tmp_path / "exports/active/in_progress/other/resume.txt"
    shared.parent.mkdir(parents=True)
    shared.write_text("other", encoding="utf-8")
    other = {
        "id": "other",
        "company": "Other",
        "role": "Other Role",
        "material_paths": {"Tailored Resume": str(shared)},
    }
    selected = _application(material_paths={"Tailored Resume": str(shared)})

    result = preflight_package_generation(
        selected["id"], [other, selected], PROJECT_ROOT, export_root=tmp_path / "exports"
    )

    assert result["status"] == "conflict"
    assert result["conflicts"][0]["conflicting_prospect_id"] == "other"


def test_gitlab_style_preflight_with_only_job_path_is_ok(tmp_path):
    record = _application(
        material_paths={
            "Job Description": str((PROJECT_ROOT / GITLAB_FIXTURE).resolve())
        }
    )
    result = preflight_package_generation(
        record["id"], [record], PROJECT_ROOT, export_root=tmp_path / "exports"
    )
    assert result["status"] == "ready"
    assert result["conflicts"] == []


def test_identical_regeneration_reuses_canonical_file_without_suffix(tmp_path):
    export_root = tmp_path / "exports"
    record = _application()
    first_source = tmp_path / "first.txt"
    first_source.write_text("same bytes", encoding="utf-8")
    first = organize_package_outputs(
        tmp_path, record, {"resume_text": str(first_source)}, export_root=export_root
    )
    current = Path(first["outputs"]["resume_text"])
    second_source = tmp_path / "second.txt"
    second_source.write_text("same bytes", encoding="utf-8")
    second = organize_package_outputs(
        tmp_path,
        record,
        {"resume_text": str(second_source)},
        preserve_existing=True,
        export_root=export_root,
    )

    assert Path(second["outputs"]["resume_text"]) == current
    assert not second_source.exists()
    assert not list(current.parent.glob("*_2.txt"))
    assert second["manifest"]["files"]["resume_text"] == str(current)
    assert Path(second["manifest"]["manifest_path"]).name == "manifest.json"


def test_changed_same_prospect_content_versions_prior_and_keeps_current_name(tmp_path):
    export_root = tmp_path / "exports"
    record = _application()
    old_source = tmp_path / "old.txt"
    old_source.write_text("old", encoding="utf-8")
    first = organize_package_outputs(
        tmp_path, record, {"resume_text": str(old_source)}, export_root=export_root
    )
    current = Path(first["outputs"]["resume_text"])
    new_source = tmp_path / "new.txt"
    new_source.write_text("new", encoding="utf-8")

    second = organize_package_outputs(
        tmp_path, record, {"resume_text": str(new_source)}, export_root=export_root
    )

    assert Path(second["outputs"]["resume_text"]) == current
    assert current.read_text(encoding="utf-8") == "new"
    versions = list((current.parent / "versions").glob("*/*.txt"))
    assert len(versions) == 1
    assert versions[0].read_text(encoding="utf-8") == "old"
    manifest = json.loads(Path(second["manifest"]["manifest_path"]).read_text())
    assert manifest["files"] == {"resume_text": str(current)}


def test_docx_regeneration_ignores_package_metadata_but_not_document_content(tmp_path):
    export_root = tmp_path / "exports"
    record = _application()
    first_source = tmp_path / "first.docx"
    first_document = Document()
    first_document.add_paragraph("Same visible résumé content")
    first_document.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    first_document.save(first_source)
    first = organize_package_outputs(
        tmp_path, record, {"ats_docx": str(first_source)}, export_root=export_root
    )
    current = Path(first["outputs"]["ats_docx"])

    second_source = tmp_path / "second.docx"
    second_document = Document()
    second_document.add_paragraph("Same visible résumé content")
    second_document.core_properties.modified = datetime(2026, 2, 1, tzinfo=timezone.utc)
    second_document.save(second_source)
    assert current.read_bytes() != second_source.read_bytes()

    second = organize_package_outputs(
        tmp_path, record, {"ats_docx": str(second_source)}, export_root=export_root
    )

    assert Path(second["outputs"]["ats_docx"]) == current
    assert not second_source.exists()
    assert not list(current.parent.glob("*_2.docx"))
    assert not (current.parent / "versions").exists()


def test_existing_manifest_for_another_prospect_blocks_install(tmp_path):
    export_root = tmp_path / "exports"
    record = _application()
    folder = export_root / "active/in_progress" / record["id"]
    folder.mkdir(parents=True)
    (folder / "manifest.json").write_text(
        json.dumps({"prospect_id": "other"}), encoding="utf-8"
    )
    source = tmp_path / "new.txt"
    source.write_text("new", encoding="utf-8")

    with pytest.raises(ValueError, match="belongs to prospect 'other'"):
        organize_package_outputs(
            tmp_path, record, {"resume_text": str(source)}, export_root=export_root
        )
    assert source.read_text(encoding="utf-8") == "new"


def test_gitlab_tailoring_is_compact_and_role_relevant():
    foundation = load_resume_foundation(PROJECT_ROOT)
    resume = _render_markdown(
        foundation,
        _parsed_gitlab_job(),
        {"top_matching_skills": []},
        "executive_operations",
    )

    assert "Career Catalyst" in resume
    assert "Microsoft Teams" in resume
    assert "Airtable" in resume
    assert "office hours" in resume
    assert "champion network" in resume.lower()
    assert "source of truth" in resume
    assert "Disney+ launch" in resume
    assert "CampaignOS" in resume
    assert "Additional Early Experience" not in resume
    assert "Trafficking Assistant" not in resume
    assert "Tylie Jones" not in resume
    assert "Professional Development" not in resume
    assert "Publishing & Creative" not in resume
    assert "OMG23 (Omnicom Media Group)" in resume
    assert public_claim_violations(resume) == []


def test_gitlab_docx_exports_keep_only_tailored_platforms(tmp_path):
    root = tmp_path / "runtime"
    (root / "data").mkdir(parents=True)
    (root / "templates/docx").mkdir(parents=True)
    (root / "exports/markdown").mkdir(parents=True)
    shutil.copy2(PROJECT_ROOT / "data/platforms.yml", root / "data/platforms.yml")
    shutil.copy2(
        PROJECT_ROOT / "templates/docx/styled_resume_template.docx",
        root / "templates/docx/styled_resume_template.docx",
    )
    foundation = load_resume_foundation(PROJECT_ROOT)
    source = root / "exports/markdown/gitlab.md"
    source.write_text(
        _render_markdown(
            foundation,
            _parsed_gitlab_job(),
            {"top_matching_skills": []},
            "executive_operations",
        ),
        encoding="utf-8",
    )

    ats = Document(export_ats_docx(source, root)["output_path"])
    styled = Document(export_styled_docx(source, root)["output_path"])
    ats_text = "\n".join(paragraph.text for paragraph in ats.paragraphs)
    styled_text = "\n".join(paragraph.text for paragraph in styled.paragraphs)
    styled_text += "\n" + "\n".join(
        cell.text for table in styled.tables for row in table.rows for cell in row.cells
    )
    for text in (ats_text, styled_text):
        assert "Microsoft Teams" in text
        assert "Airtable" in text
        assert "AI Workflow Design" in text
        assert "Publishing & Creative" not in text
        assert "Photography" not in text


def test_gitlab_cover_letter_uses_prioritized_grounded_evidence():
    context = load_generation_context(GITLAB_FIXTURE, PROJECT_ROOT)
    content = _cover_letter_content(context)

    assert 250 <= _word_count(content) <= 325
    assert content.count("Career Catalyst") >= 2
    assert "Airtable" in content
    assert "Microsoft Teams" in content
    assert "office hours" in content
    assert "champion network" in content.lower()
    assert "Disney+ launch" in content
    assert "10 direct reports" in content
    assert "64-person organization" in content
    assert "GitLab" in content
    assert "CampaignOS" in content and "working prototype" in content
    assert "more operational than flashy" not in content.lower()
    assert "I care about the work itself" not in content
    assert public_claim_violations(content) == []


def test_internal_employer_aliases_remain_available_but_public_display_is_current():
    foundation = load_resume_foundation(PROJECT_ROOT)
    aliases = historical_omg23_aliases(foundation)
    assert "OMD Entertainment" in aliases
    assert "OMG23 / OMD Entertainment" in aliases
    assert company_display_name("Gitlab") == "GitLab"
    assert company_display_name("OMD Entertainment") == "OMG23 (Omnicom Media Group)"
    assert company_display_name("OMG23 / OMD Entertainment") == "OMG23 (Omnicom Media Group)"
