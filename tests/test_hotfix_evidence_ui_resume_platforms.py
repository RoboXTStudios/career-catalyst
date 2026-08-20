from __future__ import annotations

import copy
import shutil
from pathlib import Path

import yaml

from tests.fixture_support import replace_evidence_projects
from docx import Document

import app
from scripts.export_docx import export_ats_docx, export_styled_docx
from scripts.resume_foundation import load_resume_foundation
from scripts.tailor_resume import _render_markdown


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _evidence(project_id: str, *, status: str = "Active", title: str | None = None):
    return {
        "id": project_id,
        "title": title or project_id.replace("_", " ").title(),
        "problem": "Verified problem.",
        "actions": "Verified actions.",
        "results": "Verified results.",
        "status": status,
    }


def _write_evidence(root: Path, projects: list[dict]) -> None:
    path = root / "data/evidence_projects.yml"
    path.parent.mkdir(parents=True, exist_ok=True)
    replace_evidence_projects(path, projects)


def _platform_section(markdown: str) -> str:
    return markdown.split("## Platforms & Technologies", 1)[1].split(
        "## Professional Experience", 1
    )[0]


def _document_text(path: str) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *tables])


def test_actual_relevant_evidence_options_merge_code_catalog_once(
    tmp_path: Path, monkeypatch
):
    runtime_root = tmp_path / "runtime"
    code_root = tmp_path / "code"
    _write_evidence(
        runtime_root,
        [
            _evidence("runtime_project", title="Runtime Project"),
            _evidence("archived_project", status="Archived"),
        ],
    )
    _write_evidence(
        code_root,
        [
            _evidence("enterprise_media_operations_transformation"),
            _evidence("multiverse_editorial", title="OMG23 Multiverse Newsletter"),
        ],
    )
    monkeypatch.setenv("CAREER_CATALYST_CODE_ROOT", str(code_root))

    _projects, _labels, options = app._relevant_evidence_options(runtime_root)

    assert options.count("multiverse_editorial") == 1
    assert "runtime_project" in options
    assert "enterprise_media_operations_transformation" in options
    assert "archived_project" not in options

    _write_evidence(
        runtime_root,
        [
            _evidence("runtime_project", title="Runtime Project"),
            _evidence("multiverse_editorial", title="Runtime Multiverse Record"),
        ],
    )
    projects, labels, options = app._relevant_evidence_options(runtime_root)
    assert options.count("multiverse_editorial") == 1
    assert len([item for item in projects if item["id"] == "multiverse_editorial"]) == 1
    assert labels["multiverse_editorial"].startswith("Runtime Multiverse Record")


def test_resume_payload_filters_contaminated_source_platforms():
    foundation = copy.deepcopy(load_resume_foundation(PROJECT_ROOT))
    foundation["data"]["platforms"]["platform_categories"] = [
        {
            "name": "Business Productivity & Collaboration",
            "items": ["Notion", "Substack", "Airtable", "Microsoft Teams"],
        }
    ]
    markdown = _render_markdown(
        foundation,
        {
            "job_title": "Operations Lead",
            "company": "Example",
            "raw_text": "Operations leadership",
            "keywords": ["operations"],
        },
        {"keyword_matches": [], "transferable_strengths": []},
        "executive_operations",
        complete_foundation=True,
    )
    platforms = _platform_section(markdown)

    assert "Notion" not in platforms
    assert "Substack" not in platforms
    assert "Airtable" in platforms
    assert "Microsoft Teams" in platforms


def test_job_mentions_cannot_add_excluded_resume_platforms():
    foundation = copy.deepcopy(load_resume_foundation(PROJECT_ROOT))
    productivity = next(
        category
        for category in foundation["data"]["platforms"]["platform_categories"]
        if category["name"] == "Business Productivity & Collaboration"
    )
    productivity["items"].append("Substack")
    markdown = _render_markdown(
        foundation,
        {
            "job_title": "Collaboration Operations Lead",
            "company": "Example",
            "raw_text": "Lead workflows using Notion, Substack, Airtable, and Microsoft Teams.",
            "keywords": ["Notion", "Substack", "Airtable", "Microsoft Teams"],
        },
        {"keyword_matches": [], "transferable_strengths": []},
        "executive_operations",
    )
    platforms = _platform_section(markdown)

    assert "Notion" not in platforms
    assert "Substack" not in platforms
    assert "Airtable" in platforms
    assert "Microsoft Teams" in platforms


def test_ats_and_styled_renderers_filter_excluded_platforms(tmp_path: Path):
    root = tmp_path / "render"
    (root / "data").mkdir(parents=True)
    (root / "templates/docx").mkdir(parents=True)
    (root / "exports/markdown").mkdir(parents=True)
    shutil.copy2(PROJECT_ROOT / "data/platforms.yml", root / "data/platforms.yml")
    shutil.copy2(
        PROJECT_ROOT / "templates/docx/styled_resume_template.docx",
        root / "templates/docx/styled_resume_template.docx",
    )
    source = root / "exports/markdown/contaminated.md"
    source.write_text(
        """<!-- career-catalyst-job-title: Operations Lead -->
<!-- career-catalyst-company: Example -->
# Trisha Lynch

## Platforms & Technologies

### Business Productivity & Collaboration

Notion, Substack, Airtable, Microsoft Teams

## Professional Experience

### Example

- Led operations.
""",
        encoding="utf-8",
    )

    ats_text = _document_text(export_ats_docx(source, root)["output_path"])
    styled_text = _document_text(export_styled_docx(source, root)["output_path"])
    for rendered in (ats_text, styled_text):
        assert "Notion" not in rendered
        assert "Substack" not in rendered
        assert "Airtable" in rendered
        assert "Microsoft Teams" in rendered
