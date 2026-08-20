import unittest
import shutil
import tempfile
from pathlib import Path

from docx import Document

from scripts.export_docx import export_ats_docx, export_styled_docx
from scripts.generate_application_note import generate_application_note
from scripts.generate_cover_letter import (
    ApplicationMaterialError,
    generate_cover_letter,
    load_generation_context,
    save_material,
)
from scripts.generate_messages import generate_message
from scripts.generate_strategy_pack import generate_strategy_pack
from scripts.tailor_resume import tailor_resume


PROJECT_ROOT = Path(__file__).resolve().parents[1]
GOOGLE_JOB = "jobs/google_strategy_ops_lead_youtube_auction_brand.md"
UNSUPPORTED_GOOGLE_CLAIMS = (
    "worked at Google",
    "inside Google",
    "employed by Google",
    "Google employee",
    "owned Google products",
    "partnered directly with Google executives",
    "internal referral",
    "I know the team",
    "family connection",
    "family relationship",
    "family at Google",
    "family member at Google",
    "relative at Google",
)


def _document_text(document):
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(paragraphs + table_cells)


class GoogleYouTubeCalibrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temporary.name)
        for directory in ("data", "config", "templates", "jobs"):
            shutil.copytree(PROJECT_ROOT / directory, cls.root / directory)
        cls.results = {
            "resume": tailor_resume("executive_operations", GOOGLE_JOB, cls.root),
            "cover_letter": generate_cover_letter(GOOGLE_JOB, cls.root),
            "recruiter": generate_message("recruiter", GOOGLE_JOB, cls.root),
            "hiring_manager": generate_message(
                "hiring-manager", GOOGLE_JOB, cls.root
            ),
            "application_note": generate_application_note(GOOGLE_JOB, cls.root),
            "strategy_pack": generate_strategy_pack(GOOGLE_JOB, cls.root),
        }
        cls.contents = {
            name: Path(result["output_path"]).read_text(encoding="utf-8")
            for name, result in cls.results.items()
        }

        markdown_path = cls.results["resume"]["output_path"]
        cls.styled_result = export_styled_docx(markdown_path, cls.root)
        cls.ats_result = export_ats_docx(markdown_path, cls.root)
        cls.styled_document = Document(cls.styled_result["output_path"])
        cls.ats_document = Document(cls.ats_result["output_path"])

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def test_google_materials_include_platform_familiarity(self):
        for name, content in self.contents.items():
            self.assertIn("YouTube", content, name)
            self.assertIn("Google", content, name)

    def test_google_materials_prioritize_gtm_and_product_activation(self):
        combined = "\n".join(self.contents.values()).lower()
        self.assertTrue("gtm" in combined or "go-to-market" in combined)
        self.assertIn("product activation", combined)
        self.assertIn("large advertiser", combined)

    def test_google_materials_do_not_imply_employment_access_or_family_ties(self):
        for name, content in self.contents.items():
            lowered = content.lower()
            for phrase in UNSUPPORTED_GOOGLE_CLAIMS:
                self.assertNotIn(phrase.lower(), lowered, name)

    def test_google_relationship_claim_guard_rejects_unsupported_language(self):
        context = load_generation_context(GOOGLE_JOB, self.root)

        with self.assertRaises(ApplicationMaterialError):
            save_material(
                context,
                "Recruiter_Message",
                "I worked at Google and have a family connection on the team.",
                minimum_words=1,
                maximum_words=100,
            )

    def test_campaignos_is_supporting_proof_not_the_headline(self):
        combined = "\n".join(self.contents.values())
        self.assertIn("CampaignOS", combined)
        self.assertIn("Career Catalyst", combined)
        self.assertLess(combined.find("YouTube"), combined.find("CampaignOS"))

    def test_disney_experience_remains_evidence_of_advertiser_scale(self):
        combined = "\n".join(
            self.contents[name]
            for name in ("resume", "cover_letter", "hiring_manager", "strategy_pack")
        )
        self.assertIn("multimillion-dollar", combined)
        self.assertRegex(combined.lower(), r"theatrical|streaming")
        self.assertIn("Disney+", combined)

    def test_google_docx_exports_preserve_positioning_and_table_invariants(self):
        styled_text = _document_text(self.styled_document)
        ats_text = _document_text(self.ats_document)

        self.assertGreaterEqual(len(self.styled_document.tables), 1)
        self.assertEqual(len(self.ats_document.tables), 0)
        for content in (styled_text, ats_text):
            self.assertIn("YouTube Product Activation", content)
            self.assertIn("GTM Operations", content)
            self.assertIn("Career Catalyst", content)
            self.assertNotIn("worked at Google", content)

    def test_google_outputs_use_expected_short_filenames(self):
        self.assertEqual(
            Path(self.styled_result["output_path"]).name,
            "google_strategy_operations_lead_youtube_auction_brand_trisha_lynch_styled_resume.docx",
        )
        self.assertEqual(
            Path(self.ats_result["output_path"]).name,
            "google_strategy_operations_lead_youtube_auction_brand_trisha_lynch_ats_resume.docx",
        )
        self.assertEqual(
            Path(self.results["cover_letter"]["output_path"]).name,
            "google_strategy_operations_lead_youtube_auction_brand_trisha_lynch_cover_letter.txt",
        )


if __name__ == "__main__":
    unittest.main()
