"""Export tailored Markdown resumes to styled or ATS-safe DOCX files."""

import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor, Twips

try:
    from .filename_utils import build_upload_filename, short_company_name
    from .load_data import DataLoadError, load_yaml_file
    from .parse_job import JobParseError, parse_job_description
    from .resume_platforms import filter_resume_platform_items
except ImportError:
    from filename_utils import build_upload_filename, short_company_name
    from load_data import DataLoadError, load_yaml_file
    from parse_job import JobParseError, parse_job_description
    from resume_platforms import filter_resume_platform_items


PathInput = Union[str, Path]
STYLED_MODE = "styled"
ATS_MODE = "ats"
SUPPORTED_MODES = (STYLED_MODE, ATS_MODE)
STYLED_TABLE_INDENT_DXA = 60
STYLED_TABLE_CELL_MARGINS_DXA = {"top": 40, "bottom": 40, "start": 60, "end": 60}


class DocxExportError(Exception):
    """Base exception for DOCX export errors."""


class MissingMarkdownFileError(DocxExportError):
    """Raised when the requested Markdown source file does not exist."""

    def __init__(self, file_path: PathInput) -> None:
        super().__init__(f"Markdown resume file not found: {file_path}")
        self.file_path = file_path


def _resolve_path(file_path: PathInput, project_root: Path) -> Path:
    path = Path(file_path)
    return path if path.is_absolute() else project_root / path


def _filename_key(value: str) -> str:
    return "".join(re.findall(r"[a-z0-9]+", value.lower()))


def _source_matches_company(source_path: Path, company: str) -> bool:
    source_key = _filename_key(source_path.stem)
    company_keys = (
        _filename_key(company),
        _filename_key(short_company_name(company)),
    )
    return any(key and key in source_key for key in company_keys)


def _legacy_resume_context(source_path: Path, project_root: Path) -> Dict[str, str]:
    tracker_path = project_root / "data" / "application_tracker.yml"
    if tracker_path.is_file():
        try:
            tracker = load_yaml_file("data/application_tracker.yml", project_root)
        except DataLoadError:
            tracker = {}
        for application in tracker.get("applications", []):
            if not isinstance(application, dict):
                continue
            company = str(application.get("company") or "")
            role = str(application.get("role") or "")
            if company and role and _source_matches_company(source_path, company):
                return {"company": company, "job_title": role}

    job_matches = []
    jobs_directory = project_root / "jobs"
    if jobs_directory.is_dir():
        for job_path in sorted(jobs_directory.iterdir()):
            if (
                not job_path.is_file()
                or job_path.name.lower() == "readme.md"
                or job_path.suffix.lower() not in {".md", ".txt"}
            ):
                continue
            try:
                parsed = parse_job_description(job_path)
            except JobParseError:
                continue
            company = str(parsed.get("company") or "")
            role = str(parsed.get("job_title") or "")
            if company and role and _source_matches_company(source_path, company):
                job_matches.append({"company": company, "job_title": role})
    if len(job_matches) == 1:
        return job_matches[0]
    return {"company": "Company", "job_title": "Role"}


def _resume_context(markdown: str, source_path: Path, project_root: Path) -> Dict[str, str]:
    metadata: Dict[str, str] = {}
    for line in markdown.splitlines():
        match = re.match(
            r"^<!--\s*career-catalyst-(job-title|company):\s*(.+?)\s*-->$",
            line.strip(),
            flags=re.IGNORECASE,
        )
        if match:
            metadata[match.group(1).lower().replace("-", "_")] = match.group(2).strip()

    candidate_match = re.search(r"^#\s+(.+?)\s*$", markdown, flags=re.MULTILINE)
    metadata["candidate_name"] = candidate_match.group(1).strip() if candidate_match else "Trisha Lynch"

    if not metadata.get("job_title") or not metadata.get("company"):
        legacy = _legacy_resume_context(source_path, project_root)
        metadata.setdefault("job_title", legacy["job_title"])
        metadata.setdefault("company", legacy["company"])
    return metadata


def _load_platform_categories(project_root: Path) -> List[Dict[str, Any]]:
    try:
        platforms = load_yaml_file("data/platforms.yml", project_root)
    except DataLoadError as error:
        raise DocxExportError(f"Could not load canonical platforms data: {error}") from error

    categories = platforms.get("platform_categories")
    if not isinstance(categories, list) or not categories:
        raise DocxExportError(
            "Canonical platforms data must contain at least one category."
        )

    for category in categories:
        if not isinstance(category, dict) or not isinstance(category.get("name"), str):
            raise DocxExportError("Each canonical platform category must have a name.")
        items = category.get("items")
        if not isinstance(items, list) or not all(isinstance(item, str) for item in items):
            raise DocxExportError(
                f"Canonical platform category '{category['name']}' must contain an items list."
            )
    return [
        {
            **category,
            "items": filter_resume_platform_items(category.get("items", [])),
        }
        for category in categories
    ]


def _set_style_font(
    style: Any,
    name: str,
    size: float,
    bold: Optional[bool] = None,
    italic: Optional[bool] = None,
    color: Optional[Tuple[int, int, int]] = None,
) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = RGBColor(*color)

    run_properties = style._element.get_or_add_rPr()
    run_fonts = run_properties.rFonts
    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, run_fonts)
    run_fonts.set(qn("w:ascii"), name)
    run_fonts.set(qn("w:hAnsi"), name)
    run_fonts.set(qn("w:eastAsia"), name)


def _get_or_create_paragraph_style(document: DocumentObject, name: str) -> Any:
    if name in document.styles:
        return document.styles[name]
    return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _format_run(
    run: Any,
    font_name: str,
    size: float,
    bold: bool = False,
    color: Tuple[int, int, int] = (0x33, 0x33, 0x33),
) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(*color)

    run_properties = run._element.get_or_add_rPr()
    run_fonts = run_properties.rFonts
    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, run_fonts)
    run_fonts.set(qn("w:ascii"), font_name)
    run_fonts.set(qn("w:hAnsi"), font_name)
    run_fonts.set(qn("w:eastAsia"), font_name)


def _set_paragraph_format(
    style: Any,
    before: float,
    after: float,
    line_spacing: float,
    keep_with_next: bool = False,
) -> None:
    paragraph_format = style.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph_format.line_spacing = line_spacing
    paragraph_format.keep_with_next = keep_with_next
    paragraph_format.widow_control = True


def _configure_styled_styles(document: DocumentObject) -> Dict[str, Any]:
    body_font = "Aptos"
    body_color = (0x33, 0x33, 0x33)
    heading_color = (0x1E, 0x35, 0x57)

    normal = document.styles["Normal"]
    _set_style_font(normal, body_font, 9.5, color=body_color)
    _set_paragraph_format(normal, before=0, after=2, line_spacing=1.03)

    heading_1 = document.styles["Heading 1"]
    _set_style_font(heading_1, body_font, 12, bold=True, color=heading_color)
    _set_paragraph_format(heading_1, before=6, after=2, line_spacing=1.0, keep_with_next=True)

    heading_2 = document.styles["Heading 2"]
    _set_style_font(heading_2, body_font, 9.7, bold=True, color=body_color)
    _set_paragraph_format(heading_2, before=4, after=1, line_spacing=1.0, keep_with_next=True)

    name_style = _get_or_create_paragraph_style(document, "Resume Name")
    _set_style_font(name_style, "Garamond", 24, bold=True, color=body_color)
    _set_paragraph_format(name_style, before=0, after=0, line_spacing=1.0, keep_with_next=True)

    positioning_style = _get_or_create_paragraph_style(document, "Resume Positioning")
    _set_style_font(positioning_style, body_font, 11, color=body_color)
    _set_paragraph_format(positioning_style, before=0, after=1, line_spacing=1.0, keep_with_next=True)

    contact_style = _get_or_create_paragraph_style(document, "Resume Contact")
    _set_style_font(contact_style, body_font, 8.8, color=body_color)
    _set_paragraph_format(contact_style, before=0, after=4, line_spacing=1.0)

    metadata_style = _get_or_create_paragraph_style(document, "Resume Metadata")
    _set_style_font(metadata_style, body_font, 8.8, italic=True, color=(0x55, 0x55, 0x55))
    _set_paragraph_format(metadata_style, before=0, after=1, line_spacing=1.0, keep_with_next=True)

    bullet_style = _get_or_create_paragraph_style(document, "Resume Bullet")
    _set_style_font(bullet_style, body_font, 9.5, color=body_color)
    _set_paragraph_format(bullet_style, before=0, after=0.5, line_spacing=1.02)

    competency_style = _get_or_create_paragraph_style(document, "Resume Competencies")
    _set_style_font(competency_style, body_font, 9.2, color=body_color)
    _set_paragraph_format(competency_style, before=0, after=2, line_spacing=1.0)

    platform_style = _get_or_create_paragraph_style(document, "Resume Platform Group")
    _set_style_font(platform_style, body_font, 9.0, color=body_color)
    _set_paragraph_format(platform_style, before=0, after=1, line_spacing=1.0)

    return {
        "body_font": body_font,
        "bullet_text_indent": 540,
        "bullet_hanging": 270,
        "competency_separator": " • ",
    }


def _configure_ats_styles(document: DocumentObject) -> Dict[str, Any]:
    body_font = "Calibri"
    black = (0, 0, 0)

    normal = document.styles["Normal"]
    _set_style_font(normal, body_font, 10.5, color=black)
    _set_paragraph_format(normal, before=0, after=2, line_spacing=1.05)

    heading_1 = document.styles["Heading 1"]
    _set_style_font(heading_1, body_font, 11.5, bold=True, color=black)
    _set_paragraph_format(heading_1, before=7, after=2, line_spacing=1.0, keep_with_next=True)

    heading_2 = document.styles["Heading 2"]
    _set_style_font(heading_2, body_font, 10.5, bold=True, color=black)
    _set_paragraph_format(heading_2, before=4, after=1, line_spacing=1.0, keep_with_next=True)

    name_style = _get_or_create_paragraph_style(document, "Resume Name")
    _set_style_font(name_style, body_font, 18, bold=True, color=black)
    _set_paragraph_format(name_style, before=0, after=1, line_spacing=1.0, keep_with_next=True)

    positioning_style = _get_or_create_paragraph_style(document, "Resume Positioning")
    _set_style_font(positioning_style, body_font, 10.5, bold=True, color=black)
    _set_paragraph_format(positioning_style, before=0, after=1, line_spacing=1.0, keep_with_next=True)

    contact_style = _get_or_create_paragraph_style(document, "Resume Contact")
    _set_style_font(contact_style, body_font, 10, color=black)
    _set_paragraph_format(contact_style, before=0, after=4, line_spacing=1.0)

    metadata_style = _get_or_create_paragraph_style(document, "Resume Metadata")
    _set_style_font(metadata_style, body_font, 10, italic=True, color=black)
    _set_paragraph_format(metadata_style, before=0, after=1, line_spacing=1.0, keep_with_next=True)

    bullet_style = _get_or_create_paragraph_style(document, "Resume Bullet")
    _set_style_font(bullet_style, body_font, 10.5, color=black)
    _set_paragraph_format(bullet_style, before=0, after=1.5, line_spacing=1.03)

    competency_style = _get_or_create_paragraph_style(document, "Resume Competencies")
    _set_style_font(competency_style, body_font, 10.5, color=black)
    _set_paragraph_format(competency_style, before=0, after=3, line_spacing=1.03)

    platform_style = _get_or_create_paragraph_style(document, "Resume Platform Group")
    _set_style_font(platform_style, body_font, 10.5, color=black)
    _set_paragraph_format(platform_style, before=0, after=1.5, line_spacing=1.03)

    return {
        "body_font": body_font,
        "bullet_text_indent": 720,
        "bullet_hanging": 360,
        "competency_separator": "; ",
    }


def _configure_page(document: DocumentObject, mode: str) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        if mode == STYLED_MODE:
            section.top_margin = Inches(0.45)
            section.bottom_margin = Inches(0.45)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)
            section.header_distance = Inches(0.25)
            section.footer_distance = Inches(0.25)
        else:
            section.top_margin = Inches(0.65)
            section.bottom_margin = Inches(0.65)
            section.left_margin = Inches(0.65)
            section.right_margin = Inches(0.65)
            section.header_distance = Inches(0.3)
            section.footer_distance = Inches(0.3)

        columns = section._sectPr.find(qn("w:cols"))
        if columns is None:
            columns = OxmlElement("w:cols")
            section._sectPr.append(columns)
        columns.set(qn("w:num"), "1")


def _ensure_child(parent: Any, tag: str) -> Any:
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_dxa_width(parent: Any, tag: str, width_dxa: int) -> None:
    width = _ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def _set_cell_margins(cell: Any, margins_dxa: Dict[str, int]) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    cell_margins = _ensure_child(cell_properties, "w:tcMar")
    for side in ("top", "bottom", "start", "end"):
        margin = _ensure_child(cell_margins, f"w:{side}")
        margin.set(qn("w:w"), str(margins_dxa[side]))
        margin.set(qn("w:type"), "dxa")


def _set_table_borders_none(table: Any) -> None:
    borders = _ensure_child(table._tbl.tblPr, "w:tblBorders")
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = _ensure_child(borders, f"w:{edge_name}")
        edge.set(qn("w:val"), "nil")


def _column_widths_from_weights(weights: List[int], total_width_dxa: int) -> List[int]:
    total_weight = sum(weights)
    widths = [round(total_width_dxa * weight / total_weight) for weight in weights]
    widths[-1] += total_width_dxa - sum(widths)
    return widths


def _apply_table_geometry(table: Any, column_widths_dxa: List[int]) -> None:
    table_width_dxa = sum(column_widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    table_properties = table._tbl.tblPr
    _set_dxa_width(table_properties, "w:tblW", table_width_dxa)

    table_indent = _ensure_child(table_properties, "w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(STYLED_TABLE_INDENT_DXA))

    layout = _ensure_child(table_properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    table_grid = table._tbl.tblGrid
    for child in list(table_grid):
        table_grid.remove(child)
    for width in column_widths_dxa:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        table_grid.append(grid_column)

    for column_index, width in enumerate(column_widths_dxa):
        table.columns[column_index].width = Twips(width)

    for row in table.rows:
        row.height = None
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for column_index, cell in enumerate(row.cells):
            width = column_widths_dxa[column_index]
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_dxa_width(cell._tc.get_or_add_tcPr(), "w:tcW", width)
            _set_cell_margins(cell, STYLED_TABLE_CELL_MARGINS_DXA)


def _styled_table_width_dxa(document: DocumentObject) -> int:
    section = document.sections[0]
    content_width = (
        section.page_width.twips - section.left_margin.twips - section.right_margin.twips
    )
    return int(content_width - STYLED_TABLE_INDENT_DXA)


def _next_numbering_id(numbering: Any, tag: str, attribute: str) -> int:
    values = []
    for element in numbering.findall(qn(tag)):
        value = element.get(qn(attribute))
        if value is not None:
            values.append(int(value))
    return max(values, default=0) + 1


def _create_bullet_numbering(
    document: DocumentObject,
    font_name: str,
    text_indent: int,
    hanging: int,
) -> int:
    numbering = document.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_id))

    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract_num.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)

    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "bullet")
    level.append(num_format)

    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•")
    level.append(level_text)

    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.append(level_justification)

    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(text_indent))
    tabs.append(tab)
    paragraph_properties.append(tabs)

    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), str(text_indent))
    indentation.set(qn("w:hanging"), str(hanging))
    paragraph_properties.append(indentation)
    level.append(paragraph_properties)

    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), font_name)
    run_fonts.set(qn("w:hAnsi"), font_name)
    run_properties.append(run_fonts)
    level.append(run_properties)

    abstract_num.append(level)
    numbering.append(abstract_num)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def _apply_bullet_numbering(paragraph: Any, num_id: int) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_properties.append(level)
    num_properties.append(number)
    paragraph_properties.append(num_properties)


def _clear_document_body(document: DocumentObject) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _load_document(project_root: Path, mode: str) -> Dict[str, Any]:
    template_path = project_root / "templates" / "docx" / "styled_resume_template.docx"
    if mode == STYLED_MODE and template_path.is_file():
        document = Document(str(template_path))
        _clear_document_body(document)
        return {"document": document, "template_path": str(template_path)}
    return {"document": Document(), "template_path": None}


def _parse_markdown(markdown: str) -> List[Dict[str, str]]:
    blocks: List[Dict[str, str]] = []
    paragraph_lines: List[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append({"type": "paragraph", "text": " ".join(paragraph_lines).strip()})
            paragraph_lines.clear()

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            continue
        if re.match(r"^<!--\s*career-catalyst-(?:job-title|company):", line, re.IGNORECASE):
            flush_paragraph()
            continue
        if line.startswith("# "):
            flush_paragraph()
            blocks.append({"type": "h1", "text": line[2:].strip()})
        elif line.startswith("## "):
            flush_paragraph()
            blocks.append({"type": "h2", "text": line[3:].strip()})
        elif line.startswith("### "):
            flush_paragraph()
            blocks.append({"type": "h3", "text": line[4:].strip()})
        elif re.match(r"^[-*]\s+", line):
            flush_paragraph()
            blocks.append({"type": "bullet", "text": re.sub(r"^[-*]\s+", "", line)})
        else:
            paragraph_lines.append(line)

    flush_paragraph()
    return blocks


def _inline_tokens(text: str) -> List[Dict[str, Any]]:
    pattern = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))")
    tokens: List[Dict[str, Any]] = []
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            tokens.append(
                {"text": text[position : match.start()], "bold": False, "url": None}
            )
        token = match.group(0)
        if token.startswith("**"):
            tokens.append({"text": token[2:-2], "bold": True, "url": None})
        else:
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            tokens.append(
                {
                    "text": link_match.group(1) if link_match else token,
                    "bold": False,
                    "url": link_match.group(2) if link_match else None,
                }
            )
        position = match.end()
    if position < len(text):
        tokens.append({"text": text[position:], "bold": False, "url": None})
    return tokens or [{"text": text, "bold": False, "url": None}]


def add_hyperlink(paragraph: Any, text: str, url: str) -> Any:
    """Add a visible external hyperlink to a python-docx paragraph."""
    if not re.match(r"^(?:https?://|mailto:)", url, flags=re.IGNORECASE):
        return paragraph.add_run(text)

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)

    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _add_inline_text(paragraph: Any, text: str, hyperlinks: bool = False) -> None:
    for token in _inline_tokens(text):
        if hyperlinks and token["url"]:
            add_hyperlink(paragraph, token["text"], token["url"])
            continue
        run = paragraph.add_run(token["text"])
        if token["bold"]:
            run.bold = True


def _is_section(section_name: str, expected: str) -> bool:
    return section_name.strip().lower() == expected.lower()


def _skip_platform_blocks(blocks: List[Dict[str, str]], start_index: int) -> int:
    index = start_index
    while index < len(blocks) and blocks[index]["type"] != "h2":
        index += 1
    return index


def _platform_categories_from_blocks(
    blocks: List[Dict[str, str]], fallback: List[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    """Use the tailored Markdown selection; canonical YAML is only a legacy fallback."""
    in_platforms = False
    categories: List[Dict[str, Any]] = []
    for index, block in enumerate(blocks):
        if block["type"] == "h2":
            if in_platforms:
                break
            in_platforms = _is_section(block["text"], "Platforms & Technologies")
            continue
        if not in_platforms or block["type"] != "h3":
            continue
        next_block = blocks[index + 1] if index + 1 < len(blocks) else None
        if not next_block or next_block["type"] != "paragraph":
            continue
        items = filter_resume_platform_items(next_block["text"].split(","))
        if items:
            categories.append({"name": block["text"], "items": items})
    return categories or fallback


def _add_styled_platforms_table(
    document: DocumentObject,
    platform_categories: List[Dict[str, Any]],
) -> None:
    column_count = min(3, len(platform_categories))
    row_count = (len(platform_categories) + column_count - 1) // column_count
    table = document.add_table(rows=row_count, cols=column_count)
    table.style = "Normal Table"
    _set_table_borders_none(table)

    cells = [cell for row in table.rows for cell in row.cells]
    for cell, category in zip(cells, platform_categories):
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_together = True

        heading_run = paragraph.add_run(category["name"])
        _format_run(heading_run, "Aptos", 7.8, bold=True)
        heading_run.add_break()

        tools_run = paragraph.add_run(", ".join(category["items"]))
        _format_run(tools_run, "Aptos", 7.5)

    column_widths = _column_widths_from_weights(
        [1] * column_count,
        _styled_table_width_dxa(document),
    )
    _apply_table_geometry(table, column_widths)


def _add_ats_platform_groups(
    document: DocumentObject,
    platform_categories: List[Dict[str, Any]],
) -> None:
    for category in platform_categories:
        paragraph = document.add_paragraph(style="Resume Platform Group")
        label_run = paragraph.add_run(f"{category['name']}: ")
        label_run.bold = True
        _add_inline_text(paragraph, ", ".join(category["items"]))


def _add_markdown_content(
    document: DocumentObject,
    blocks: List[Dict[str, str]],
    bullet_num_id: int,
    competency_separator: str,
    mode: str,
    platform_categories: List[Dict[str, Any]],
) -> None:
    def add_text(paragraph: Any, text: str) -> None:
        _add_inline_text(paragraph, text, hyperlinks=mode == STYLED_MODE)

    header_paragraph_index = 0
    current_section = ""
    current_subheading = ""
    subheading_paragraph_index = 0
    index = 0

    while index < len(blocks):
        block = blocks[index]
        block_type = block["type"]
        text = block["text"]

        if block_type == "h1":
            paragraph = document.add_paragraph(style="Resume Name")
            add_text(paragraph, text)
            index += 1
            continue

        if block_type == "h2":
            current_section = text
            current_subheading = ""
            subheading_paragraph_index = 0
            paragraph = document.add_paragraph(style="Heading 1")
            add_text(paragraph, text)
            index += 1
            if _is_section(text, "Platforms & Technologies"):
                index = _skip_platform_blocks(blocks, index)
                if mode == STYLED_MODE:
                    _add_styled_platforms_table(document, platform_categories)
                else:
                    _add_ats_platform_groups(document, platform_categories)
            continue

        if block_type == "h3":
            current_subheading = text
            subheading_paragraph_index = 0
            next_block = blocks[index + 1] if index + 1 < len(blocks) else None
            if (
                _is_section(current_section, "Platforms & Technologies")
                and next_block
                and next_block["type"] == "paragraph"
            ):
                paragraph = document.add_paragraph(style="Resume Platform Group")
                label_run = paragraph.add_run(f"{text}: ")
                label_run.bold = True
                add_text(paragraph, next_block["text"])
                index += 2
                continue

            paragraph = document.add_paragraph(style="Heading 2")
            add_text(paragraph, text)
            index += 1
            continue

        if block_type == "bullet" and _is_section(current_section, "Core Competencies"):
            competencies = []
            while index < len(blocks) and blocks[index]["type"] == "bullet":
                competencies.append(blocks[index]["text"])
                index += 1
            paragraph = document.add_paragraph(style="Resume Competencies")
            add_text(paragraph, competency_separator.join(competencies))
            continue

        if block_type == "bullet":
            paragraph = document.add_paragraph(style="Resume Bullet")
            _apply_bullet_numbering(paragraph, bullet_num_id)
            add_text(paragraph, text)
            index += 1
            continue

        if not current_section:
            style = "Resume Positioning" if header_paragraph_index == 0 else "Resume Contact"
            header_paragraph_index += 1
        elif current_subheading and subheading_paragraph_index == 0:
            style = "Resume Metadata"
            subheading_paragraph_index += 1
        else:
            style = "Normal"
            if current_subheading:
                subheading_paragraph_index += 1

        paragraph = document.add_paragraph(style=style)
        add_text(paragraph, text)
        index += 1


def _output_path(
    markdown: str,
    source_path: Path,
    project_root: Path,
    mode: str,
) -> Path:
    context = _resume_context(markdown, source_path, project_root)
    suffix = "Styled" if mode == STYLED_MODE else "ATS"
    filename = build_upload_filename(
        context["candidate_name"],
        context["job_title"],
        context["company"],
        suffix,
        "docx",
    )
    return project_root / "exports" / "docx" / filename


def _export_docx(
    markdown_path: PathInput,
    mode: str,
    project_root: Optional[PathInput] = None,
) -> Dict[str, Any]:
    if mode not in SUPPORTED_MODES:
        raise DocxExportError(f"Unsupported DOCX export mode: {mode}")

    root = Path(project_root) if project_root is not None else Path.cwd()
    source_path = _resolve_path(markdown_path, root)
    if not source_path.is_file():
        raise MissingMarkdownFileError(markdown_path)

    try:
        markdown = source_path.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as error:
        raise DocxExportError(f"Unable to read Markdown resume {markdown_path}: {error}") from error

    document_info = _load_document(root, mode)
    document = document_info["document"]
    blocks = _parse_markdown(markdown)
    platform_categories = _platform_categories_from_blocks(
        blocks, _load_platform_categories(root)
    )
    _configure_page(document, mode)
    style_config = (
        _configure_styled_styles(document)
        if mode == STYLED_MODE
        else _configure_ats_styles(document)
    )
    bullet_num_id = _create_bullet_numbering(
        document,
        style_config["body_font"],
        style_config["bullet_text_indent"],
        style_config["bullet_hanging"],
    )
    _add_markdown_content(
        document,
        blocks,
        bullet_num_id,
        style_config["competency_separator"],
        mode,
        platform_categories,
    )

    output_path = _output_path(markdown, source_path, root, mode)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        document.save(str(output_path))
    except OSError as error:
        raise DocxExportError(f"Unable to save DOCX resume {output_path}: {error}") from error

    return {
        "mode": mode,
        "source_path": str(source_path),
        "output_path": str(output_path),
        "template_path": document_info["template_path"],
    }


def export_styled_docx(
    markdown_path: PathInput,
    project_root: Optional[PathInput] = None,
) -> Dict[str, Any]:
    """Export a polished DOCX, using the styled template when available."""
    return _export_docx(markdown_path, STYLED_MODE, project_root)


def export_ats_docx(
    markdown_path: PathInput,
    project_root: Optional[PathInput] = None,
) -> Dict[str, Any]:
    """Export a plain, single-column, no-table ATS-safe DOCX."""
    return _export_docx(markdown_path, ATS_MODE, project_root)


def export_docx(
    markdown_path: PathInput,
    project_root: Optional[PathInput] = None,
) -> Dict[str, Any]:
    """Export a styled DOCX for backward compatibility with Sprint 5."""
    return export_styled_docx(markdown_path, project_root)
